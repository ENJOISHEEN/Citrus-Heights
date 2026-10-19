"""Huddle board pipeline.
  python3 board.py seed  <stack_ranking.xlsx> <state.json>          # exact MTD from a full file (re-sync)
  python3 board.py update <state.json> <dump.txt> <nps_email.txt> <data_date YYYY-MM-DD>   # add one day's daily sheets
  python3 board.py resync <state.json> <store_mtd.tsv> <employee_mtd.tsv>   # exact MTD from rows pasted out of the workbook (Store/Employee MTD tabs)
  python3 build_board.py <state.json> <publish_date YYYY-MM-DD> <out.html>
"""
import sys,json,re,datetime as dt
from dumpparse import parse_dump
from nps_parse import parse_nps,CODES
KEYS=['rgu','int','tv','voice','home','mt','mo','sel','plus','leg','wear','acc']
STORE='59010'
def rank_of(state):
    DAY=dt.date.fromisoformat(state['date']).day
    t=sorted(state['market'].items(),key=lambda kv:-kv[1]['rgu']/(kv[1]['target']*DAY/31))
    return [c for c,_ in t].index(STORE)+1
def seed(xlsx,out):
    import openpyxl
    wb=openpyxl.load_workbook(xlsx,data_only=True)
    st={}
    def f(x): return float(x or 0)
    market={}
    for r in wb['Store-MTD'].iter_rows(values_only=True):
        m=re.match(r'\[(\d+)\]',str(r[1] or ''))
        if m and m.group(1) in CODES:
            market[m.group(1)]=dict(name=CODES[m.group(1)],target=int(r[5]),rgu=int(r[6]),traffic=int(r[10]),int=int(r[12]),tv=int(r[14]),voice=int(r[17]),home=int(r[19]),mt=int(r[21]),mo=int(r[23]),sel=int(r[25]),plus=int(r[27]),leg=int(r[29]),wear=int(f(r[43])),acc=round(f(r[44]),2))
    daily={}
    for r in wb['Employee-Daily'].iter_rows(values_only=True):
        if r[3] and f'[{STORE}]' in str(r[3]): daily[r[1]]=int(f(r[6]))
    reps=[]
    for r in wb['Employee-MTD'].iter_rows(values_only=True):
        if r[3] and f'[{STORE}]' in str(r[3]):
            reps.append(dict(name=r[1],rgu=int(r[7]),int=int(r[11]),tv=int(r[12]),voice=int(r[13]),home=int(r[14]),mt=int(r[15]),mo=int(r[16]),sel=int(r[17]),plus=int(r[19]),leg=int(r[21]),wear=int(f(r[34])),acc=round(f(r[35]),2),daily=daily.get(r[1],0)))
    reps.sort(key=lambda x:-x['rgu'])
    sdr=0
    for r in wb['Store-Daily'].iter_rows(values_only=True):
        if r[1] and f'[{STORE}]' in str(r[1]): sdr=int(f(r[5]))
    date=None
    for r in wb['Store-MTD'].iter_rows(values_only=True):
        if r[0] and 'Date Range' in str(r[0]) or (len(r)>1 and r[1] and 'Date Range' in str(r[1])): pass
    ds=[c for row in wb['Store-MTD'].iter_rows(max_row=8,values_only=True) for c in row if c and 'Date Range' in str(c)][0]
    m=re.search(r'- (\d+)/(\d+)/(\d+)',ds); date=f'{m.group(3)}-{int(m.group(1)):02d}-{int(m.group(2)):02d}'
    s=market[STORE]
    state=dict(date=date,target=s['target'],plan_history=st.get('plan_history',[740,750,761,724,s['target']]),store={k:s[k] for k in ['rgu','int','tv','voice','home','mt','mo','sel','plus','leg','wear','acc','traffic']},reps=reps,market=market,nps={c:list(v) for c,v in (parse_nps(open(sys.argv[4]).read()) if len(sys.argv)>4 else {}).items()},daily_store_rgu=sdr,prev_rank=None)
    state['prev_rank']=rank_of(state)
    json.dump(state,open(out,'w'),indent=1); print('seeded',date,s['rgu'],len(reps),'reps')
def update(sp,dump,npsf,ddate):
    st=json.load(open(sp)); new=dt.date.fromisoformat(ddate)
    if dt.date.fromisoformat(st['date'])>=new: sys.exit(f"state already at {st['date']}; refusing to double-count {ddate}")
    if (new-dt.date.fromisoformat(st['date'])).days!=1: sys.exit(f"gap: state {st['date']} -> {ddate}; re-sync from a full file first")
    sh=parse_dump(open(dump).read())
    rep_daily={}; sdaily={}
    for r in sh['Employee-Daily']:
        if len(r)>19 and r[3] and f'[{STORE}]' in str(r[3]): rep_daily[r[1]]=r
    for r in sh['Store-Daily']:
        m=re.match(r'\[(\d+)\]',str(r[1] if len(r)>1 else ''))
        if m and m.group(1) in CODES: sdaily[m.group(1)]=r
    if STORE not in sdaily or not rep_daily: sys.exit('daily sheets missing Sunrise rows')
    z=lambda x:x if isinstance(x,(int,float)) else 0
    st['prev_rank']=rank_of(st)
    for c,r in sdaily.items():
        s=st['market'][c]
        for k,i in (('rgu',5),('traffic',6),('int',8),('tv',10),('voice',12),('home',14),('mt',16),('mo',18),('sel',20),('plus',22),('leg',24),('wear',29),('acc',30)): s[k]=round(s[k]+z(r[i]),2) if k=='acc' else int(s[k]+z(r[i]))
    s=st['market'][STORE]; st['store']={**st['store'],**{k:s[k] for k in ['rgu','int','tv','voice','home','mt','mo','sel','plus','leg','wear','acc','traffic']}}
    st['daily_store_rgu']=int(z(sdaily[STORE][5]))
    seen=set(); notes=[]
    for r in st['reps']:
        d=rep_daily.get(r['name']); seen.add(r['name'])
        r['daily']=int(z(d[6])) if d else 0
        if d:
            for k,i in (('rgu',6),('int',7),('tv',8),('voice',9),('home',10),('mt',11),('mo',12),('sel',13),('plus',14),('leg',15),('wear',18),('acc',19)): r[k]=round(r[k]+z(d[i]),2) if k=='acc' else int(r[k]+z(d[i]))
    for name,d in rep_daily.items():
        if name not in seen:
            nr=dict(name=name,daily=int(z(d[6])),**{k:(round(z(d[i]),2) if k=='acc' else int(z(d[i]))) for k,i in (('rgu',6),('int',7),('tv',8),('voice',9),('home',10),('mt',11),('mo',12),('sel',13),('plus',14),('leg',15),('wear',18),('acc',19))})
            st['reps'].append(nr); notes.append('NEW REP '+name+' (MTD starts at today)')
    st['reps'].sort(key=lambda x:-x['rgu'])
    chk=sum(r['rgu'] for r in st['reps'])-st['store']['rgu']
    if chk: notes.append(f'rep RGU sum differs from store by {chk}')
    notes.append('targets/rep totals not verified since last resync; paste MTD rows and run board.py resync')
    nps=parse_nps(open(npsf).read())
    if STORE not in nps: sys.exit('NPS email missing Sunrise')
    st['nps']={c:list(v) for c,v in nps.items()}
    st['date']=ddate
    json.dump(st,open(sp,'w'),indent=1); print('updated to',ddate,'store RGU',st['store']['rgu'],'|',notes)
def resync(sp,store_tsv,emp_tsv):
    """Overwrite running totals with exact MTD numbers pasted from the workbook (tab-separated rows, no header).
    Store rows: Store, RGU Target, Activated RGU, Trend%, Traffic, Conv%, Internet, IntConv%, TV, Voice, Home, Mobile&Tablet, M&T Conv%, Mobile Only, MO Conv%, Select, Sel%, Plus, Plus%, Legacy, Leg%, Wearables, Accessory.
    Employee rows: Employee, RGU, Trend%, Internet, TV, Voice, Home, Mobile&Tablet, Mobile Only, Select, Sel%, Plus, Plus%, Legacy, Leg%, Wearables, Accessory."""
    st=json.load(open(sp))
    n=lambda x:float(str(x).replace('$','').replace(',','').replace('%','').strip())
    SK=['target','rgu',None,'traffic',None,'int',None,'tv','voice','home','mt',None,'mo',None,'sel',None,'plus',None,'leg',None,'wear','acc']
    SK2=['target','rgu',None,'traffic',None,'int',None,'tv','voice',None,'home',None,'mt',None,'mo',None,'sel',None,'plus',None,'leg',None,'wear','acc']  # newer layout adds Voice/Home conversion columns
    RK=['rgu',None,'int','tv','voice','home','mt','mo','sel',None,'plus',None,'leg',None,'wear','acc']
    ch=[]
    for l in open(store_tsv):
        c=l.rstrip('\r\n').split('\t'); m=re.match(r'\[(\d+)\]',c[0])
        if not m or m.group(1) not in st['market']: continue
        mk=st['market'][m.group(1)]
        for k,v in zip(SK2 if len(c)-1==len(SK2) else SK,c[1:]):
            if k:
                v=round(n(v),2) if k=='acc' else int(n(v))
                if mk[k]!=v: ch.append(f"{mk['name']} {k} {mk[k]}->{v}"); mk[k]=v
    s=st['market'][STORE]
    if s['target']!=st['target']:
        st['target']=s['target']
        if st['plan_history'][-1]!=s['target']: st['plan_history'].append(s['target'])
    st['store']={**st['store'],**{k:s[k] for k in st['store']}}
    reps={r['name']:r for r in st['reps']}
    for l in open(emp_tsv):
        c=l.rstrip('\r\n').split('\t')
        if c[0] not in reps: ch.append('NEW REP '+c[0]); reps[c[0]]=dict(name=c[0],daily=0); st['reps'].append(reps[c[0]])
        r=reps[c[0]]
        for k,v in zip(RK,c[1:]):
            if k:
                v=round(n(v),2) if k=='acc' else int(n(v))
                if r.get(k)!=v: ch.append(f"{c[0]} {k} {r.get(k)}->{v}"); r[k]=v
    st['reps'].sort(key=lambda x:-x['rgu'])
    chk=sum(r['rgu'] for r in st['reps'])-st['store']['rgu']
    json.dump(st,open(sp,'w'),indent=1); print('resynced',st['date'],len(ch),'changes'); print('\n'.join(ch)); print('rep RGU sum vs store:',chk)
if __name__=='__main__':
    if sys.argv[1]=='seed': seed(sys.argv[2],sys.argv[3])
    elif sys.argv[1]=='resync': resync(*sys.argv[2:5])
    elif sys.argv[1]=='update': update(*sys.argv[2:6])
